"""Builders for PDF and DOCX test documents (no binary fixtures in the repository)."""

import io


def make_pdf(pages: list[list[str]]) -> bytes:
    """A minimal text-layer PDF: one list of lines per page, Helvetica 11pt."""
    objects: list[bytes] = []
    n_pages = len(pages)
    font_id = 3 + 2 * n_pages
    kids = " ".join(f"{3 + 2 * i} 0 R" for i in range(n_pages))
    objects.append(b"<< /Type /Catalog /Pages 2 0 R >>")
    objects.append(f"<< /Type /Pages /Kids [{kids}] /Count {n_pages} >>".encode())
    for i, lines in enumerate(pages):
        content_id = 4 + 2 * i
        objects.append(
            f"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 595 842] /Contents {content_id} 0 R "
            f"/Resources << /Font << /F1 {font_id} 0 R >> >> >>".encode()
        )
        ops = ["BT", "/F1 11 Tf", "14 TL", "50 790 Td"]
        for line in lines:
            safe = line.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")
            ops.append(f"({safe}) Tj T*")
        ops.append("ET")
        stream = "\n".join(ops).encode("latin-1")
        objects.append(b"<< /Length %d >>\nstream\n" % len(stream) + stream + b"\nendstream")
    objects.append(b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>")
    out = io.BytesIO()
    out.write(b"%PDF-1.4\n")
    offsets = []
    for i, obj in enumerate(objects, start=1):
        offsets.append(out.tell())
        out.write(f"{i} 0 obj\n".encode() + obj + b"\nendobj\n")
    xref = out.tell()
    out.write(f"xref\n0 {len(objects) + 1}\n0000000000 65535 f \n".encode())
    for off in offsets:
        out.write(f"{off:010d} 00000 n \n".encode())
    out.write(f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF\n".encode())
    return out.getvalue()


def make_docx(blocks: list[tuple[str, str]]) -> bytes:
    """blocks: ("h1"|"h2"|"p", text) or ("table", "a|b;c|d")."""
    import docx

    document = docx.Document()
    for kind, text in blocks:
        if kind.startswith("h"):
            document.add_heading(text, int(kind[1]))
        elif kind == "table":
            rows = [r.split("|") for r in text.split(";")]
            table = document.add_table(rows=len(rows), cols=len(rows[0]))
            for r, row in enumerate(rows):
                for c, value in enumerate(row):
                    table.cell(r, c).text = value
        else:
            document.add_paragraph(text)
    buf = io.BytesIO()
    document.save(buf)
    return buf.getvalue()
