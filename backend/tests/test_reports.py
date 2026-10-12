"""Investigation report in Markdown, Word and PDF: same content, audited with the file's checksum."""

import hashlib
import io

from docx import Document as Docx
from pypdf import PdfReader
from sqlalchemy import select

from app.models import AuditEvent
from app.services import report_render
from tests.conftest import auth
from tests.test_api import _new_incident


def _report(client, inc_id: str, fmt: str, who: str = "auditor"):
    return client.get(f"/incidents/{inc_id}/report", params={"format": fmt}, headers=auth(client, who))


def test_word_and_pdf_carry_the_same_report(client, project_id, db):
    inc = _new_incident(client, project_id)
    client.post(f"/incidents/{inc['id']}/investigate", headers=auth(client, "safety"))
    md = _report(client, inc["id"], "md")
    assert md.status_code == 200 and md.headers["content-type"].startswith("text/markdown")

    docx = _report(client, inc["id"], "docx")
    assert docx.status_code == 200 and docx.headers["content-type"].startswith(
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    )
    assert f"{inc['reference']}-report.docx" in docx.headers["content-disposition"]
    assert docx.headers["X-Report-SHA256"] == hashlib.sha256(docx.content).hexdigest()
    word = Docx(io.BytesIO(docx.content))
    text = "\n".join(p.text for p in word.paragraphs)
    cells = "\n".join(c.text for t in word.tables for row in t.rows for c in row.cells)
    assert text.startswith(f"Investigation report {inc['reference']}")
    assert "AI investigation (AI-generated, reviewed by people)" in text
    assert "not a legal compliance determination" in text and "Not yet closable" in text
    assert "PENDING_REVIEW" in cells and "**" not in text + cells

    pdf = _report(client, inc["id"], "pdf")
    assert pdf.status_code == 200 and pdf.headers["content-type"] == "application/pdf"
    assert pdf.content.startswith(b"%PDF-")
    pages = PdfReader(io.BytesIO(pdf.content)).pages
    extracted = " ".join(" ".join(p.extract_text().split()) for p in pages)
    assert f"Investigation report {inc['reference']}" in extracted
    assert "AI-generated, reviewed by people" in extracted and "PENDING_REVIEW" in extracted

    events = db.scalars(
        select(AuditEvent).where(AuditEvent.entity_id == inc["id"], AuditEvent.action == "incident.report")
    ).all()
    assert {e.details["format"] for e in events} == {"md", "docx", "pdf"}
    assert pdf.headers["X-Report-SHA256"] in {e.details["sha256"] for e in events}


def test_report_format_and_scope_are_checked(client, project_id):
    inc = _new_incident(client, project_id)
    bad = _report(client, inc["id"], "html")
    assert bad.status_code == 400 and bad.json()["error"]["code"] == "INVALID_REQUEST"
    assert _report(client, inc["id"], "pdf", who="other").status_code == 404
    assert client.get(f"/incidents/{inc['id']}/report", params={"format": "docx"}).status_code == 401


def test_renderer_keeps_structure_and_handles_unicode():
    md = "\n".join(
        [
            "# Report SG-1: Fall — scaffold",
            "",
            "> Decision support only.",
            "",
            "| Field | Value |",
            "|---|---|",
            "| Risk | **HIGH** ≥ 15 |",
            "| Empty |  |",
            "",
            "1. Why one",
            "- Unsafe act: no harness",
            "_Synthetic data._",
            "**Triage:** fall from height; Zürich site 区域",
        ]
    )
    blocks = report_render.parse(md)
    assert [b.kind for b in blocks] == ["h1", "note", "table", "li", "li", "note", "p"]
    assert blocks[2].rows == [["Field", "Value"], ["Risk", "**HIGH** ≥ 15"], ["Empty", ""]]
    word = Docx(io.BytesIO(report_render.render(md, "docx", "t")))
    assert "Zürich site 区域" in "\n".join(p.text for p in word.paragraphs)  # Word keeps every character
    pdf = report_render.render(md, "pdf", "t")
    text = " ".join(p.extract_text() for p in PdfReader(io.BytesIO(pdf)).pages)
    assert "Fall - scaffold" in text and ">= 15" in text and "Zürich" in text  # Latin-1 subset in the PDF
