"""Incident reports (Playbook Prompt 34).

A report is built only from stored records: the incident, its history, evidence, the latest AI run outputs, the
risk assessment, CAPA actions with every approval decision and verification, and the sources those records cite.
Nothing is generated at report time. AI output appears in its own clearly labelled sections with the agent,
provider and run id, so a reader can tell recorded facts from model suggestions.

`build` returns a list of blocks; `to_markdown`, `to_docx` and `to_pdf` render the same blocks, so the three formats
always carry the same content. Rendering happens in this service only (no user-supplied templates or HTML).
"""

import hashlib
import io
import re
from dataclasses import dataclass
from datetime import datetime

from sqlalchemy import desc, or_, select
from sqlalchemy.orm import Session

from app.models import (
    AgentRun,
    ApprovalEvent,
    CapaAction,
    Document,
    DocumentChunk,
    Evidence,
    Incident,
    IncidentEvent,
    Project,
    RiskAssessment,
    Site,
    User,
    VerificationRecord,
    utcnow,
)

DISCLAIMER = (
    "Site Guard AI is decision support. Sections marked AI-GENERATED are model suggestions that people reviewed or "
    "must review; they are not verified facts. Nothing in this report is a legal or regulatory compliance "
    "determination or a certification."
)
FORMATS = {
    "md": "text/markdown; charset=utf-8",
    "docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    "pdf": "application/pdf",
}
_REF = re.compile(r"\b(chunk|evidence):([0-9a-f-]{36})\b")


@dataclass
class Block:
    kind: str  # h1 | h2 | p | note | ai | table
    text: str = ""
    headers: tuple[str, ...] = ()
    rows: tuple[tuple[str, ...], ...] = ()


def _fmt(value) -> str:
    if value is None or value == "":
        return "-"
    if isinstance(value, datetime):
        return value.strftime("%Y-%m-%d %H:%M UTC")
    return str(getattr(value, "value", value))


def _refs_in(obj) -> list[str]:
    found: list[str] = []
    if isinstance(obj, dict):
        for v in obj.values():
            found += _refs_in(v)
    elif isinstance(obj, list):
        for v in obj:
            found += _refs_in(v)
    elif isinstance(obj, str):
        found += [f"{kind}:{rid}" for kind, rid in _REF.findall(obj)]
    return found


def _list(items) -> str:
    return "; ".join(str(i) for i in items) if items else "-"


def build(db: Session, inc: Incident, generated_at: datetime | None = None, *, include_ai: bool = True) -> list[Block]:
    generated_at = generated_at or utcnow()
    project = db.get(Project, inc.project_id)
    site = db.get(Site, inc.site_id) if inc.site_id else None
    names: dict[str, str] = {}

    def who(user_id: str | None) -> str:
        if not user_id:
            return "-"
        if user_id not in names:
            user = db.get(User, user_id)
            names[user_id] = user.full_name if user else "unknown user"
        return names[user_id]

    blocks: list[Block] = [
        Block("h1", f"Incident report {inc.reference}: {inc.title}"),
        Block(
            "p",
            f"Project {project.code if project else '-'} {project.name if project else ''}. "
            f"Generated {_fmt(generated_at)} from the records held in Site Guard AI.",
        ),
        Block("note", DISCLAIMER),
    ]
    if inc.is_synthetic:
        blocks.append(Block("note", "This incident is synthetic demo data."))

    blocks += [
        Block("h2", "1. Incident record"),
        Block(
            "table",
            headers=("Field", "Value"),
            rows=(
                ("Reference", inc.reference),
                ("Status", _fmt(inc.status)),
                ("Domain", _fmt(inc.domain)),
                ("Reported severity", _fmt(inc.severity)),
                ("Category", _fmt(inc.category)),
                ("Occurred", _fmt(inc.occurred_at)),
                ("Reported", _fmt(inc.reported_at)),
                ("Reported by", who(inc.reporter_id)),
                ("Site", site.name if site else "-"),
                ("Location", _fmt(inc.location)),
                ("Activity", _fmt(inc.activity)),
                ("People involved", str(inc.people_involved)),
                ("Closed", _fmt(inc.closed_at)),
            ),
        ),
        Block("p", f"Description: {inc.description}"),
        Block("p", f"Immediate actions: {_fmt(inc.immediate_actions)}"),
    ]

    events = db.scalars(
        select(IncidentEvent).where(IncidentEvent.incident_id == inc.id).order_by(IncidentEvent.created_at)
    )
    blocks += [
        Block("h2", "2. Status history"),
        Block(
            "table",
            headers=("When", "From", "To", "By", "Note"),
            rows=tuple(
                (_fmt(e.created_at), _fmt(e.from_status), e.to_status, who(e.actor_id), _fmt(e.note)) for e in events
            ),
        ),
    ]

    evidence = list(db.scalars(select(Evidence).where(Evidence.incident_id == inc.id).order_by(Evidence.created_at)))
    blocks += [
        Block("h2", "3. Evidence"),
        Block(
            "table",
            headers=("File", "Type", "SHA-256", "Uploaded by", "When", "Description"),
            rows=tuple(
                (
                    e.filename,
                    e.content_type,
                    e.sha256[:16] + "...",
                    who(e.uploaded_by),
                    _fmt(e.created_at),
                    _fmt(e.description),
                )
                for e in evidence
            ),
        )
        if evidence
        else Block("p", "No evidence files are attached."),
    ]

    risk = db.scalar(
        select(RiskAssessment)
        .where(RiskAssessment.incident_id == inc.id)
        .order_by(desc(RiskAssessment.created_at))
        .limit(1)
    )
    blocks.append(Block("h2", "4. Risk assessment"))
    if risk:
        source = (
            "suggested by AI, scored by the deterministic matrix"
            if risk.inputs_source == "AI_SUGGESTED"
            else "assessed by a person"
        )
        blocks.append(
            Block(
                "p",
                f"Matrix {risk.matrix_version}: likelihood {risk.likelihood} x consequence {risk.consequence} = "
                f"{risk.score} ({risk.band}). Inputs {source}; recorded by {who(risk.assessed_by)} on "
                f"{_fmt(risk.created_at)}. {risk.rationale}",
            )
        )
    else:
        blocks.append(Block("p", "No risk assessment is recorded."))

    runs = list(db.scalars(select(AgentRun).where(AgentRun.incident_id == inc.id).order_by(AgentRun.started_at)))
    latest = [r for r in runs if runs and r.workflow_id == runs[-1].workflow_id]
    blocks.append(Block("h2", "5. AI analysis (AI-GENERATED, for human review)"))
    cited: list[str] = []
    if not include_ai:
        latest = []
        blocks.append(Block("p", "Left out of this report on request."))
    elif not latest:
        blocks.append(Block("p", "No AI analysis has been run."))
    for r in latest:
        label = f"{r.agent} agent, provider {r.provider}, run {r.id}"
        if r.status != "SUCCEEDED" or not r.output:
            blocks.append(Block("ai", f"[{label}] Failed ({_fmt(r.error)}); no output was used."))
            continue
        out = r.output
        cited += _refs_in(out)
        review = " Flagged for human review." if r.needs_human_review else ""
        blocks.append(
            Block("ai", f"[{label}] Confidence {out.get('confidence', '-')}.{review} {_summary(r.agent, out)}")
        )
        if out.get("open_questions"):
            blocks.append(Block("ai", f"[{r.agent}] Open questions: {_list(out['open_questions'])}"))

    actions = list(
        db.scalars(select(CapaAction).where(CapaAction.incident_id == inc.id).order_by(CapaAction.created_at))
    )
    ids = [a.id for a in actions]
    approvals = (
        list(db.scalars(select(ApprovalEvent).where(ApprovalEvent.capa_id.in_(ids)).order_by(ApprovalEvent.created_at)))
        if ids
        else []
    )
    verifications = (
        list(
            db.scalars(
                select(VerificationRecord)
                .where(VerificationRecord.capa_id.in_(ids))
                .order_by(VerificationRecord.created_at)
            )
        )
        if ids
        else []
    )
    title_of = {a.id: a.title for a in actions}
    blocks += [
        Block("h2", "6. Corrective and preventive actions"),
        Block(
            "table",
            headers=("Action", "Type", "Origin", "Approval", "Work", "Assignee", "Due", "Critical"),
            rows=tuple(
                (
                    a.title,
                    a.action_type,
                    "AI-proposed" if a.ai_generated else "Human",
                    _fmt(a.approval_status),
                    _fmt(a.work_status),
                    who(a.assignee_id),
                    _fmt(a.due_date),
                    "yes" if a.critical else "no",
                )
                for a in actions
            ),
        )
        if actions
        else Block("p", "No actions are recorded."),
        Block("h2", "7. Approval history"),
        Block(
            "table",
            headers=("When", "Action", "Decision", "Reviewer", "Reason"),
            rows=tuple(
                (_fmt(e.created_at), title_of.get(e.capa_id, e.capa_id), e.decision, who(e.reviewer_id), e.reason)
                for e in approvals
            ),
        )
        if approvals
        else Block("p", "No review decisions are recorded."),
        Block("h2", "8. Verification"),
        Block(
            "table",
            headers=("When", "Action", "Effective", "Verifier", "Notes"),
            rows=tuple(
                (
                    _fmt(v.created_at),
                    title_of.get(v.capa_id, v.capa_id),
                    "yes" if v.effective else "no",
                    who(v.verifier_id),
                    v.notes,
                )
                for v in verifications
            ),
        )
        if verifications
        else Block("p", "No verification is recorded."),
    ]
    for a in actions:
        cited += [r for r in (a.evidence_refs or []) if isinstance(r, str)]

    blocks += [Block("h2", "9. Sources cited")] + _citations(db, inc, cited)
    return blocks


def _summary(agent: str, out: dict) -> str:
    if agent == "triage":
        return (
            f"Class {out.get('incident_class')}; severity candidate {out.get('severity_candidate')}; priority "
            f"{out.get('priority')}. Hazards: {_list(out.get('hazards'))}. {out.get('rationale', '')}"
        )
    if agent in ("safety", "quality"):
        extra = (
            f"Failed or missing controls: {_list(out.get('failed_or_missing_controls'))}."
            if agent == "safety"
            else f"Defects: {_list(out.get('defects'))}. Probable stage: {out.get('probable_stage')}."
        )
        scores = f"Suggested likelihood {out.get('likelihood')}, consequence {out.get('consequence')}."
        return f"{out.get('summary', '')} {extra} {scores}"
    if agent == "rca":
        causes = "; ".join(f"{c['category']}: {c['statement']}" for c in out.get("root_causes", []))
        return f"Problem: {out.get('problem_statement')}. Root causes: {causes}."
    if agent == "compliance":
        findings = "; ".join(
            f"{f['status']} - {f['requirement']} ({f['source_title']})" for f in out.get("findings", [])
        )
        return f"Possible gaps for review: {findings or '-'}. {out.get('disclaimer', '')}"
    if agent == "capa":
        return f"Proposed {len(out.get('actions', []))} action(s); see section 6 for the human decisions."
    return ""


def _citations(db: Session, inc: Incident, refs: list[str]) -> list[Block]:
    unique = list(dict.fromkeys(refs))
    rows: list[tuple[str, ...]] = []
    for ref in unique:
        kind, _, rid = ref.partition(":")
        if kind == "chunk":
            row = db.execute(
                select(DocumentChunk, Document)
                .join(Document, Document.id == DocumentChunk.document_id)
                .where(
                    DocumentChunk.id == rid,
                    Document.organization_id == inc.organization_id,
                    or_(Document.project_id.is_(None), Document.project_id == inc.project_id),
                )
            ).first()
            if row:
                chunk, doc = row
                where = f"section {chunk.section}" + (f", page {chunk.page}" if chunk.page else "")
                rows.append((ref, f"{doc.title} (version {doc.version})", where, doc.source))
        elif kind == "evidence":
            ev = db.get(Evidence, rid)
            if ev is not None and ev.incident_id == inc.id:
                rows.append((ref, ev.filename, f"SHA-256 {ev.sha256[:16]}...", "incident evidence"))
    if not rows:
        return [Block("p", "No document or evidence sources are cited.")]
    return [Block("table", headers=("Ref", "Source", "Location", "Origin"), rows=tuple(rows))]


# --- renderers ----------------------------------------------------------------------------------------------------


def _md_cell(text: str) -> str:
    return text.replace("|", "\\|").replace("\n", " ")


def to_markdown(blocks: list[Block]) -> str:
    out: list[str] = []
    for b in blocks:
        if b.kind == "h1":
            out.append(f"# {b.text}")
        elif b.kind == "h2":
            out.append(f"## {b.text}")
        elif b.kind == "note":
            out.append(f"> {b.text}")
        elif b.kind == "ai":
            out.append(f"> **AI-GENERATED** {b.text}")
        elif b.kind == "table":
            out.append("| " + " | ".join(b.headers) + " |")
            out.append("|" + "---|" * len(b.headers))
            out += ["| " + " | ".join(_md_cell(c) for c in row) + " |" for row in b.rows]
        else:
            out.append(b.text)
        out.append("")
    return "\n".join(out)


def to_docx(blocks: list[Block]) -> bytes:
    from docx import Document as Docx

    doc = Docx()
    for b in blocks:
        if b.kind == "h1":
            doc.add_heading(b.text, level=1)
        elif b.kind == "h2":
            doc.add_heading(b.text, level=2)
        elif b.kind == "table":
            table = doc.add_table(rows=1, cols=len(b.headers))
            table.style = "Table Grid"
            for cell, header in zip(table.rows[0].cells, b.headers, strict=True):
                cell.text = header
            for row in b.rows:
                for cell, value in zip(table.add_row().cells, row, strict=True):
                    cell.text = value
        else:
            para = doc.add_paragraph()
            if b.kind == "ai":
                para.add_run("AI-GENERATED ").bold = True
            run = para.add_run(b.text)
            run.italic = b.kind in ("note", "ai")
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


_ASCII = str.maketrans(
    {
        "\u2018": "'",
        "\u2019": "'",
        "\u201c": '"',
        "\u201d": '"',
        "\u2013": "-",
        "\u2014": "-",
        "\u2022": "*",
        "\u2265": ">=",
        "\u2264": "<=",
        "\u00d7": "x",
        "\u2192": "->",
    }
)


def _latin1(text: str) -> str:
    return text.translate(_ASCII).encode("latin-1", "replace").decode("latin-1")


def to_pdf(blocks: list[Block]) -> bytes:
    from fpdf import FPDF
    from fpdf.enums import XPos, YPos

    pdf = FPDF(format="A4")
    pdf.set_auto_page_break(auto=True, margin=15)
    pdf.set_title("Site Guard AI incident report")
    pdf.set_creator("Site Guard AI")
    pdf.add_page()
    width = pdf.w - pdf.l_margin - pdf.r_margin
    for b in blocks:
        if b.kind in ("h1", "h2"):
            pdf.set_font("Helvetica", "B", 15 if b.kind == "h1" else 12)
            pdf.multi_cell(width, 7, _latin1(b.text), new_x=XPos.LMARGIN, new_y=YPos.NEXT)
            pdf.ln(1)
        elif b.kind == "table":
            pdf.set_font("Helvetica", "", 8)
            with pdf.table(col_widths=None, text_align="LEFT", line_height=4.5) as table:
                header = table.row()
                for h in b.headers:
                    header.cell(_latin1(h))
                for row in b.rows:
                    r = table.row()
                    for c in row:
                        r.cell(_latin1(c))
            pdf.ln(2)
        else:
            style = "I" if b.kind in ("note", "ai") else ""
            pdf.set_font("Helvetica", style, 9)
            prefix = "AI-GENERATED " if b.kind == "ai" else ""
            pdf.multi_cell(width, 5, _latin1(prefix + b.text), new_x=XPos.LMARGIN, new_y=YPos.NEXT)
            pdf.ln(1)
    return bytes(pdf.output())


def render(blocks: list[Block], fmt: str) -> bytes:
    if fmt == "md":
        return to_markdown(blocks).encode()
    if fmt == "docx":
        return to_docx(blocks)
    if fmt == "pdf":
        return to_pdf(blocks)
    raise ValueError(fmt)


def checksum(content: bytes) -> str:
    return hashlib.sha256(content).hexdigest()
