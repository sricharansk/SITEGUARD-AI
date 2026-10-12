"""Word and PDF versions of the investigation report.

The Markdown from `report.build` is the single source: it is split into blocks (headings, paragraphs, notes, lists,
tables) and the same blocks are rendered to DOCX or PDF, so every format carries the same content and the AI and
human labels survive. Only this fixed subset of Markdown is produced by `report.build`; nothing is rendered as HTML
and no user-supplied template is used. Incident text inside the report is rendered as text.
"""

import hashlib
import io
import re
from dataclasses import dataclass, field

FORMATS = {
    "md": "text/markdown; charset=utf-8",
    "docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    "pdf": "application/pdf",
}
_BOLD = re.compile(r"\*\*(.+?)\*\*")
_ORDERED = re.compile(r"^\d+\.\s+")


@dataclass
class Block:
    kind: str  # h1 | h2 | p | note | li | table
    text: str = ""
    rows: list[list[str]] = field(default_factory=list)  # table: first row is the header


def _cells(line: str) -> list[str]:
    return [c.strip() for c in line.strip().strip("|").split("|")]


def parse(markdown: str) -> list[Block]:
    blocks: list[Block] = []
    lines = markdown.splitlines()
    i = 0
    while i < len(lines):
        line = lines[i].rstrip()
        if not line.strip():
            i += 1
            continue
        if line.startswith("|"):
            rows = []
            while i < len(lines) and lines[i].startswith("|"):
                if not re.fullmatch(r"\|(-+\|)+", lines[i].strip()):  # skip the |---| separator
                    rows.append(_cells(lines[i]))
                i += 1
            blocks.append(Block("table", rows=rows))
            continue
        if line.startswith("## "):
            blocks.append(Block("h2", line[3:]))
        elif line.startswith("# "):
            blocks.append(Block("h1", line[2:]))
        elif line.startswith("> "):
            blocks.append(Block("note", line[2:]))
        elif line.startswith("- ") or _ORDERED.match(line):
            blocks.append(Block("li", line))
        elif line.startswith("_") and line.endswith("_") and len(line) > 2:
            blocks.append(Block("note", line[1:-1]))
        else:
            blocks.append(Block("p", line))
        i += 1
    return blocks


def _runs(text: str) -> list[tuple[str, bool]]:
    """Split `**bold**` spans into (text, bold) runs."""
    out: list[tuple[str, bool]] = []
    pos = 0
    for m in _BOLD.finditer(text):
        if m.start() > pos:
            out.append((text[pos : m.start()], False))
        out.append((m.group(1), True))
        pos = m.end()
    if pos < len(text):
        out.append((text[pos:], False))
    return out


def _plain(text: str) -> str:
    return _BOLD.sub(r"\1", text)


def to_docx(blocks: list[Block], title: str) -> bytes:
    from docx import Document as Docx

    doc = Docx()
    doc.core_properties.title = title
    doc.core_properties.author = "Site Guard AI"
    for b in blocks:
        if b.kind in ("h1", "h2"):
            doc.add_heading(_plain(b.text), level=1 if b.kind == "h1" else 2)
        elif b.kind == "table":
            width = max(len(r) for r in b.rows)
            table = doc.add_table(rows=0, cols=width)
            table.style = "Table Grid"
            for n, row in enumerate(b.rows):
                cells = table.add_row().cells
                for cell, value in zip(cells, row + [""] * (width - len(row)), strict=True):
                    cell.text = ""
                    cell.paragraphs[0].add_run(_plain(value)).bold = n == 0
        else:
            para = doc.add_paragraph(style="List Bullet" if b.kind == "li" and b.text.startswith("- ") else None)
            text = b.text[2:] if b.kind == "li" and b.text.startswith("- ") else b.text
            for chunk, bold in _runs(text):
                run = para.add_run(chunk)
                run.bold = bold
                run.italic = b.kind == "note"
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


# The PDF uses the standard Helvetica font, which only covers Latin-1; common typographic characters are mapped and
# anything else becomes "?". Word and Markdown keep the full text.
_LATIN1 = str.maketrans(
    {
        "‘": "'",
        "’": "'",
        "“": '"',
        "”": '"',
        "–": "-",
        "—": "-",
        "•": "*",
        "…": "...",
        "≥": ">=",
        "≤": "<=",
        "→": "->",
    }
)


def _latin1(text: str) -> str:
    return _plain(text).translate(_LATIN1).encode("latin-1", "replace").decode("latin-1")


def to_pdf(blocks: list[Block], title: str) -> bytes:
    from fpdf import FPDF
    from fpdf.enums import XPos, YPos

    pdf = FPDF(format="A4")
    pdf.set_auto_page_break(auto=True, margin=15)
    pdf.set_title(_latin1(title))
    pdf.set_creator("Site Guard AI")
    pdf.add_page()
    width = pdf.w - pdf.l_margin - pdf.r_margin
    for b in blocks:
        if b.kind in ("h1", "h2"):
            pdf.set_font("Helvetica", "B", 14 if b.kind == "h1" else 12)
            pdf.multi_cell(width, 7, _latin1(b.text), new_x=XPos.LMARGIN, new_y=YPos.NEXT)
            pdf.ln(1)
        elif b.kind == "table":
            pdf.set_font("Helvetica", "", 7.5)
            cols = max(len(r) for r in b.rows)
            with pdf.table(text_align="LEFT", line_height=4, first_row_as_headings=True) as table:
                for row in b.rows:
                    r = table.row()
                    for value in row + [""] * (cols - len(row)):
                        r.cell(_latin1(value))
            pdf.ln(2)
        else:
            pdf.set_font("Helvetica", "I" if b.kind == "note" else "", 9)
            text = ("* " + b.text[2:]) if b.kind == "li" and b.text.startswith("- ") else b.text
            pdf.multi_cell(width, 4.8, _latin1(text), new_x=XPos.LMARGIN, new_y=YPos.NEXT)
            pdf.ln(1)
    return bytes(pdf.output())


def render(markdown: str, fmt: str, title: str) -> bytes:
    if fmt == "md":
        return markdown.encode()
    blocks = parse(markdown)
    if fmt == "docx":
        return to_docx(blocks, title)
    if fmt == "pdf":
        return to_pdf(blocks, title)
    raise ValueError(f"unsupported report format {fmt!r}")


def checksum(content: bytes) -> str:
    return hashlib.sha256(content).hexdigest()
