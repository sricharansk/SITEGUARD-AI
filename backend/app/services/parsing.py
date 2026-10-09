"""Document parsers (Playbook Prompt 09).

Each parser turns uploaded bytes into pages of text plus metadata, without trusting the content: text is data and
is never executed or interpreted as instructions. Parsing failures raise ParseError with a safe message; the
caller records the document as FAILED and keeps its provenance (file name, size, SHA-256).
"""

import io
import logging
import re
import zipfile
from dataclasses import dataclass, field
from pathlib import PurePath

from app.services.isolation import WorkerCrashed, WorkerTimeout, run_worker

log = logging.getLogger("siteguard.parsing")

MAX_PAGES = 500
MAX_TEXT_CHARS = 5_000_000  # extracted text per document
MAX_UNZIPPED_BYTES = 100 * 1024 * 1024  # DOCX: total uncompressed size
MAX_ZIP_RATIO = 100  # DOCX: uncompressed / compressed
MAX_ZIP_ENTRIES = 5_000


class ParseError(Exception):
    pass


@dataclass
class ParsedPage:
    number: int | None  # 1-based page number; None for formats without pages (Markdown, text, DOCX)
    text: str


@dataclass
class ParsedDocument:
    parser: str
    pages: list[ParsedPage]
    metadata: dict = field(default_factory=dict)

    @property
    def page_count(self) -> int | None:
        numbered = [p.number for p in self.pages if p.number is not None]
        return max(numbered) if numbered else None

    @property
    def text(self) -> str:
        return "\n\n".join(p.text for p in self.pages)


class Parser:
    name: str
    extensions: tuple[str, ...]
    content_types: tuple[str, ...]
    isolated = False  # binary formats are parsed in a separate, resource-limited process

    def parse(self, content: bytes) -> ParsedDocument:  # pragma: no cover - interface
        raise NotImplementedError


class TextParser(Parser):
    """Markdown and plain text. Headings (#) are kept for heading-aware chunking."""

    name = "text"
    extensions = (".md", ".markdown", ".txt")
    content_types = ("text/markdown", "text/plain", "text/x-markdown")

    def parse(self, content: bytes) -> ParsedDocument:
        try:
            text = content.decode("utf-8")
        except UnicodeDecodeError as exc:
            raise ParseError("Text documents must be UTF-8") from exc
        if not text.strip():
            raise ParseError("Document is empty")
        return ParsedDocument(self.name, [ParsedPage(None, text)])


class PdfParser(Parser):
    """Text-layer PDFs, one ParsedPage per page so chunks keep page numbers. Scanned PDFs need OCR (not built)."""

    name = "pdf"
    extensions = (".pdf",)
    content_types = ("application/pdf",)
    isolated = True

    def parse(self, content: bytes) -> ParsedDocument:
        from pypdf import PdfReader
        from pypdf.errors import PdfReadError

        if not content.startswith(b"%PDF-"):
            raise ParseError("File is not a PDF")
        try:
            reader = PdfReader(io.BytesIO(content))
            if reader.is_encrypted:
                raise ParseError("Encrypted PDFs are not supported")
            if len(reader.pages) > MAX_PAGES:
                raise ParseError(f"PDF has more than {MAX_PAGES} pages")
            pages: list[ParsedPage] = []
            total = 0
            for i, page in enumerate(reader.pages):
                text = _clean(page.extract_text() or "")
                total += len(text)
                if total > MAX_TEXT_CHARS:
                    raise ParseError(f"PDF has more than {MAX_TEXT_CHARS} characters of text")
                pages.append(ParsedPage(i + 1, text))
            info: dict = dict(reader.metadata or {})
        except ParseError:
            raise
        except PdfReadError as exc:
            raise ParseError("PDF could not be read") from exc
        except Exception as exc:  # malformed files raise many error types; none may escape as a 500
            raise ParseError("PDF could not be read") from exc
        if not any(p.text.strip() for p in pages):
            raise ParseError("PDF has no text layer (scanned PDFs need OCR, which is not available)")
        meta = {k: str(info.get(f"/{k}")) for k in ("Title", "Author", "CreationDate") if info.get(f"/{k}")}
        return ParsedDocument(self.name, pages, meta)


class DocxParser(Parser):
    """Word documents. Heading styles become Markdown headings; tables are kept as pipe-separated rows."""

    name = "docx"
    extensions = (".docx",)
    content_types = ("application/vnd.openxmlformats-officedocument.wordprocessingml.document",)
    isolated = True

    def parse(self, content: bytes) -> ParsedDocument:
        if not content.startswith(b"PK"):
            raise ParseError("File is not a DOCX document")
        _check_zip(content)
        try:
            return self._parse(content)
        except ParseError:
            raise
        except Exception as exc:  # malformed files raise many error types; none may escape as a 500
            raise ParseError("DOCX could not be read") from exc

    def _parse(self, content: bytes) -> ParsedDocument:
        import docx
        from docx.table import Table

        document = docx.Document(io.BytesIO(content))

        lines: list[str] = []
        for block in document.iter_inner_content():  # paragraphs and tables in document order
            if isinstance(block, Table):
                for row in block.rows:
                    lines.append("| " + " | ".join(c.text.strip() for c in row.cells) + " |")
                lines.append("")
                continue
            text = block.text.strip()
            if not text:
                continue
            style = (block.style.name if block.style is not None else "") or ""
            m = re.match(r"Heading (\d)", style)
            level = min(int(m.group(1)), 4) if m else (1 if style == "Title" else 0)
            lines.append(f"{'#' * level} {text}" if level else text)
            lines.append("")
        text = "\n".join(lines).strip()
        if not text:
            raise ParseError("Document is empty")
        if len(text) > MAX_TEXT_CHARS:
            raise ParseError(f"DOCX has more than {MAX_TEXT_CHARS} characters of text")
        props = document.core_properties
        meta = {k: str(v) for k, v in (("Title", props.title), ("Author", props.author)) if v}
        return ParsedDocument(self.name, [ParsedPage(None, text)], meta)


PARSERS: list[Parser] = [TextParser(), PdfParser(), DocxParser()]


def parser_for(filename: str) -> Parser:
    ext = PurePath(filename).suffix.lower()
    for parser in PARSERS:
        if ext in parser.extensions:
            return parser
    allowed = ", ".join(e for p in PARSERS for e in p.extensions)
    raise ParseError(f"Unsupported document type {ext or '(none)'}; allowed: {allowed}")


def _check_zip(content: bytes) -> None:
    """Reject zip bombs before any XML is parsed."""
    try:
        with zipfile.ZipFile(io.BytesIO(content)) as zf:
            entries = zf.infolist()
    except zipfile.BadZipFile as exc:
        raise ParseError("DOCX could not be read") from exc
    total = sum(e.file_size for e in entries)
    if len(entries) > MAX_ZIP_ENTRIES or total > MAX_UNZIPPED_BYTES or total > MAX_ZIP_RATIO * max(len(content), 1):
        raise ParseError("DOCX expands to an unsafe size")


def parse_document(filename: str, content: bytes, *, timeout: float = 60, memory_mb: int = 1024) -> ParsedDocument:
    """Pick the parser for the file name and parse. Binary formats run in a child process with CPU, memory and
    wall-clock limits, so a hostile file cannot stall or exhaust the API process."""
    parser = parser_for(filename)
    if not parser.isolated:
        return parser.parse(content)

    try:
        result = run_worker("app.services.parse_worker", [parser.name], content, timeout=timeout, memory_mb=memory_mb)
    except WorkerTimeout as exc:
        raise ParseError("Document took too long to parse") from exc
    except WorkerCrashed:
        raise ParseError("Document could not be parsed within the resource limits") from None
    if not result.get("ok"):
        raise ParseError(result.get("error") or "Document could not be read")
    pages = [ParsedPage(n, text) for n, text in result["pages"]]
    return ParsedDocument(result["parser"], pages, result.get("metadata") or {})


def _clean(text: str) -> str:
    # Join words hyphenated across line breaks and drop trailing spaces; keep line structure for headings/tables.
    text = re.sub(r"(\w)-\n(\w)", r"\1\2", text)
    return "\n".join(line.rstrip() for line in text.splitlines()).strip()
