"""Build the Review III deck: fill the programme template (stage.pptx, template slides already duplicated to 18) with
Site Guard AI content. See README.md in this folder.

    python build.py                       # stage.pptx -> ../SiteGuardAI_Review3_S.K.Sricharan_T9984.pptx
    python build.py stage.pptx out.pptx   # explicit input and output
"""

import copy
import os
import sys
from pathlib import Path

from lxml import etree
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION, XL_LEGEND_POSITION
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

HERE = Path(__file__).resolve().parent
SRC = (
    str(Path(sys.argv[1]).resolve()) if len(sys.argv) > 1 else str(HERE / "stage.pptx")
)
OUT = (
    str(Path(sys.argv[2]).resolve())
    if len(sys.argv) > 2
    else str(HERE.parent / "SiteGuardAI_Review3_S.K.Sricharan_T9984.pptx")
)
os.chdir(HERE)  # images are referenced as img/...
PROJECT = "Site Guard AI: Multi-Agent Construction Safety & Quality Incident Resolution System"

NAVY = "1F4E79"  # template title colour (5B9BD5 at 50%)
BLUE = "2E75B6"
SKY = "1BA1D6"
TEAL = "0F8B8D"
GREEN = "4C9A3C"
INK = "1F2937"
MUTED = "4B5563"
PALE = "EEF5FB"  # card fill
PALE_G = "EDF6EA"
PALE_T = "E6F4F4"
WHITE = "FFFFFF"
GREY = "9CA3AF"

prs = Presentation(SRC)
ALL = list(prs.slides)
assert len(ALL) == 18, len(ALL)
SCREENS = ALL[12:14]  # two product-screen slides after the end-to-end findings
S = ALL[:12] + ALL[14:]


def rgb(h):
    return RGBColor.from_string(h)


def strip_animation(slide):
    t = slide._element.find(qn("p:timing"))
    if t is not None:
        t.getparent().remove(t)


def remove_body_placeholder(slide):
    for sh in list(slide.shapes):
        if sh.is_placeholder:
            sh._element.getparent().remove(sh._element)


def set_runs(shape, text, size=None, color=None):
    p = shape.text_frame.paragraphs[0]
    runs = p.runs
    runs[0].text = text
    for r in runs[1:]:
        r._r.getparent().remove(r._r)
    if size:
        runs[0].font.size = Pt(size)
    if color:
        runs[0].font.color.rgb = rgb(color)
    for extra in shape.text_frame.paragraphs[1:]:
        extra._p.getparent().remove(extra._p)


def frame(slide, title, notes=None):
    """Header project title, slide title; remove the template's sample body and its click animation."""
    strip_animation(slide)
    remove_body_placeholder(slide)
    for sh in slide.shapes:
        if not sh.has_text_frame:
            continue
        t = sh.text_frame.text
        if t.startswith("Project Title"):
            sh.left, sh.width = Inches(4.35), Inches(8.85)
            set_runs(sh, PROJECT, size=13)
            sh.text_frame.paragraphs[0].runs[0].font.name = "Calibri"
        elif "Calibri body -36" in t:
            set_runs(sh, title, size=32)
    if notes:
        slide.notes_slide.notes_text_frame.text = notes


def no_shadow(shape):
    sppr = shape._element.spPr
    if sppr.find(qn("a:effectLst")) is None:
        sppr.append(etree.SubElement(sppr, qn("a:effectLst")))


def box(
    slide,
    x,
    y,
    w,
    h,
    fill=PALE,
    line=None,
    shape=MSO_SHAPE.ROUNDED_RECTANGLE,
    radius=0.08,
    name=None,
):
    s = slide.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    if shape == MSO_SHAPE.ROUNDED_RECTANGLE:
        s.adjustments[0] = radius
    if fill:
        s.fill.solid()
        s.fill.fore_color.rgb = rgb(fill)
    else:
        s.fill.background()
    if line:
        s.line.color.rgb = rgb(line)
        s.line.width = Pt(1.25)
    else:
        s.line.fill.background()
    no_shadow(s)
    if name:
        s.name = name
    s.text_frame.text = ""
    return s


def text(
    slide,
    x,
    y,
    w,
    h,
    paras,
    align=PP_ALIGN.LEFT,
    anchor=MSO_ANCHOR.TOP,
    shape=None,
    margin=0.08,
    name=None,
):
    """paras: list of (text, size, bold, color) or (runs list) where runs = [(text, size, bold, color)]."""
    if shape is None:
        shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = shape.text_frame
    tf.word_wrap = True
    tf.auto_size = None
    tf.vertical_anchor = anchor
    for side in ("left", "right", "top", "bottom"):
        setattr(tf, f"margin_{side}", Inches(margin))
    first = True
    for para in paras:
        bullet = False
        space = 4
        if isinstance(para, dict):
            bullet = para.get("bullet", False)
            space = para.get("space", 4)
            runs = para["runs"]
        elif isinstance(para[0], tuple):
            runs = para
        else:
            runs = [para]
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        p.alignment = align
        p.space_after = Pt(space)
        if bullet:
            pPr = p._p.get_or_add_pPr()
            pPr.set("marL", str(Inches(0.22)))
            pPr.set("indent", str(-Inches(0.2)))
            bu = etree.SubElement(pPr, qn("a:buChar"))
            bu.set("char", "•")
        for t, size, bold, color in runs:
            r = p.add_run()
            r.text = t
            r.font.size = Pt(size)
            r.font.bold = bold
            r.font.name = "Calibri"
            r.font.color.rgb = rgb(color)
    if name:
        shape.name = name
    return shape


def card(
    slide,
    x,
    y,
    w,
    h,
    head,
    body,
    fill=PALE,
    head_color=NAVY,
    size=14,
    head_size=16,
    badge=None,
    badge_fill=BLUE,
):
    b = box(slide, x, y, w, h, fill=fill)
    left = x
    if badge:
        c = box(
            slide, x + 0.15, y + 0.15, 0.45, 0.45, fill=badge_fill, shape=MSO_SHAPE.OVAL
        )
        text(
            slide,
            0,
            0,
            0,
            0,
            [(badge, 14, True, WHITE)],
            align=PP_ALIGN.CENTER,
            anchor=MSO_ANCHOR.MIDDLE,
            shape=c,
            margin=0,
        )
        left = x + 0.55
    text(
        slide,
        left + 0.1,
        y + 0.1,
        w - (left - x) - 0.2,
        0.5,
        [(head, head_size, True, head_color)],
        anchor=MSO_ANCHOR.MIDDLE,
    )
    paras = [
        (t, size, False, INK) if isinstance(t, str) else t
        for t in ([body] if isinstance(body, str) else body)
    ]
    off = 0.15 + head_size / 40
    text(slide, x + 0.15, y + off, w - 0.3, h - off - 0.05, paras)
    return b


def arrow(slide, x1, y1, x2, y2, color=GREY, width=1.75):
    c = slide.shapes.add_connector(
        MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2)
    )
    c.line.color.rgb = rgb(color)
    c.line.width = Pt(width)
    ln = c.line._get_or_add_ln()
    tail = etree.SubElement(ln, qn("a:tailEnd"))
    tail.set("type", "triangle")
    tail.set("w", "med")
    tail.set("len", "med")
    return c


def node(
    slide,
    x,
    y,
    w,
    h,
    title,
    sub,
    fill,
    title_color=WHITE,
    sub_color=WHITE,
    tsize=14,
    ssize=11,
):
    b = box(slide, x, y, w, h, fill=fill, radius=0.12)
    paras = [(title, tsize, True, title_color)]
    if sub:
        paras.append((sub, ssize, False, sub_color))
    text(
        slide,
        0,
        0,
        0,
        0,
        paras,
        align=PP_ALIGN.CENTER,
        anchor=MSO_ANCHOR.MIDDLE,
        shape=b,
        margin=0.05,
    )
    return b


# ---------------------------------------------------------------- 1. Title
s = S[0]
strip_animation(s)
for sh in s.shapes:
    if sh.has_text_frame and sh.text_frame.text.startswith("Review III"):
        tb = sh
paras = tb.text_frame.paragraphs
# paragraphs: 0 Review III, 1 Project Title, 2 "Calibri body -36", 3 blank, 4 Team Member, 5 "Calibri body -24", 6/7 blank
paras[1].runs[0].text = "Site Guard AI: Multi-Agent Construction Safety"
paras[1].runs[0].font.size = Pt(32)
p2 = paras[2]
p2.runs[0].text = "& Quality Incident Resolution System"
p2.runs[0].font.size = Pt(32)
for r in p2.runs[1:]:
    r._r.getparent().remove(r._r)
for idx in (3, 6, 7):  # shrink spacer paragraphs
    epr = paras[idx]._p.find(qn("a:endParaRPr"))
    if epr is not None:
        epr.set("sz", "1400")
paras[4].runs[0].text = "Team Member"
paras[4].runs[0].font.size = Pt(22)
paras[5].runs[0].text = "S.K. Sricharan"
paras[5].runs[0].font.size = Pt(26)
id_p = copy.deepcopy(paras[5]._p)
paras[5]._p.addnext(id_p)
id_runs = id_p.findall(qn("a:r"))
id_runs[0].find(qn("a:t")).text = "T-Series ID: T9984"
id_runs[0].find(qn("a:rPr")).set("sz", "2400")
s.notes_slide.notes_text_frame.text = (
    "Review III: Site Guard AI, a multi-agent system that takes a construction safety or quality incident from "
    "report to verified closure with AI decision support and human approval. Presented by S.K. Sricharan (T9984)."
)

# ---------------------------------------------------------------- 2. Introduction
s = S[1]
frame(s, "Introduction", "Context, what the system is, and the live product.")
text(
    s,
    0.45,
    1.45,
    6.85,
    4.3,
    [
        {
            "runs": [
                (
                    "Construction sites report two kinds of events that must be investigated and closed: ",
                    16,
                    False,
                    INK,
                ),
                ("safety incidents", 16, True, NAVY),
                (
                    " (falls, struck-by, excavation collapse, electrical, lifting, "
                    "plant) and ",
                    16,
                    False,
                    INK,
                ),
                ("quality defects", 16, True, NAVY),
                (
                    " (cracking, honeycombing, low cover, tolerance, water ingress).",
                    16,
                    False,
                    INK,
                ),
            ],
            "space": 10,
        },
        {
            "runs": [
                ("Site Guard AI", 16, True, NAVY),
                (
                    " is an agentic decision-support and workflow platform:",
                    16,
                    False,
                    INK,
                ),
            ],
            "space": 6,
        },
        {
            "runs": [
                (
                    "Six bounded AI agents triage and investigate each incident, find root causes, map it to company "
                    "procedures and propose corrective and preventive actions (CAPA)",
                    15,
                    False,
                    INK,
                )
            ],
            "bullet": True,
        },
        {
            "runs": [
                (
                    "Every finding cites the procedure it came from; risk is scored by a deterministic 5×5 matrix",
                    15,
                    False,
                    INK,
                )
            ],
            "bullet": True,
        },
        {
            "runs": [
                (
                    "Qualified people approve, verify and close; the AI never approves, closes or certifies "
                    "compliance",
                    15,
                    False,
                    INK,
                )
            ],
            "bullet": True,
        },
    ],
)
s.shapes.add_picture(
    "img/dashboard_after.png", Inches(7.6), Inches(1.5), width=Inches(5.3)
)
text(
    s,
    7.6,
    4.5,
    5.3,
    0.3,
    [
        (
            "Live Site Guard AI dashboard after resolving every incident (synthetic demo project)",
            11,
            False,
            MUTED,
        )
    ],
)
for i, (big, small) in enumerate(
    (
        ("20", "hazard & defect classes"),
        ("6 + 1", "AI agents + risk engine"),
        ("9", "user roles (RBAC)"),
    )
):
    x = 7.6 + i * 1.8
    b = box(s, x, 5.0, 1.65, 1.2, fill=PALE)
    text(
        s,
        0,
        0,
        0,
        0,
        [(big, 26, True, NAVY), (small, 11, False, MUTED)],
        align=PP_ALIGN.CENTER,
        anchor=MSO_ANCHOR.MIDDLE,
        shape=b,
        margin=0.04,
    )

# ---------------------------------------------------------------- 3. Objective
s = S[2]
frame(
    s, "Objectives", "Six objectives; each is delivered and measured later in the deck."
)
objs = [
    (
        "Unified incident capture",
        "Report safety and quality incidents with photos and documents, scoped by "
        "organisation and project, with a full lifecycle and audit trail.",
    ),
    (
        "AI-assisted investigation",
        "Triage, safety / quality investigation, 5-Whys root cause and compliance "
        "mapping by bounded agents with validated, structured outputs.",
    ),
    (
        "Grounded in procedures",
        "Every requirement and answer cites the retrieved company procedure; unsupported "
        "claims are removed and injected instructions are flagged.",
    ),
    (
        "Deterministic risk",
        "A 5×5 likelihood × consequence matrix computes the risk band; the AI may only suggest "
        "the inputs.",
    ),
    (
        "Human-in-the-loop CAPA",
        "AI proposes actions; authorised people approve, verify and close. No safety-critical "
        "incident is ever closed automatically.",
    ),
    (
        "Measured & deployable",
        "≥ 85% class and ≥ 90% domain accuracy on a held-out set; secure API, CI on every "
        "change, Docker image and Azure templates.",
    ),
]
for i, (h, b) in enumerate(objs):
    col, row = i % 3, i // 3
    card(
        s,
        0.45 + col * 4.17,
        1.5 + row * 2.15,
        3.97,
        1.95,
        h,
        b,
        badge=str(i + 1),
        badge_fill=[BLUE, SKY, TEAL, NAVY, GREEN, BLUE][i],
    )

# ---------------------------------------------------------------- 4. Modules overview
s = S[3]
frame(
    s, "Modules – System Overview", "Eight modules built and integrated in the product."
)
mods = [
    (
        "Incident & Evidence",
        "Reporting, lifecycle state machine, evidence files with SHA-256 integrity",
        BLUE,
    ),
    (
        "Multi-Agent Investigation",
        "Orchestrator with six agents, checkpoints, pause / resume, conflict records",
        NAVY,
    ),
    (
        "Knowledge & Retrieval",
        "PDF / DOCX upload, chunking, embeddings, hybrid search, grounded answers",
        TEAL,
    ),
    (
        "Risk & Compliance",
        "5×5 risk matrix, requirement mapping marked “possible gap, needs review”",
        SKY,
    ),
    (
        "CAPA Workflow",
        "Approval by role, overdue escalation, independent verification, closure rules",
        GREEN,
    ),
    (
        "Vision Evidence",
        "Photo checks for PPE / hazards / defects in an isolated process; engineer review",
        BLUE,
    ),
    (
        "Knowledge Graph",
        "Optional Neo4j projection of incidents, causes and actions with provenance",
        TEAL,
    ),
    (
        "Reports, Alerts & Audit",
        "Markdown / Word / PDF reports, email outbox, dashboard, full audit trail",
        NAVY,
    ),
]
for i, (h, b, c) in enumerate(mods):
    col, row = i % 4, i // 4
    x, y = 0.45 + col * 3.13, 1.5 + row * 2.15
    box(s, x, y, 2.98, 1.95, fill=PALE)
    hb = box(s, x, y, 2.98, 0.62, fill=c, radius=0.15)
    text(
        s,
        0,
        0,
        0,
        0,
        [(h, 15, True, WHITE)],
        align=PP_ALIGN.CENTER,
        anchor=MSO_ANCHOR.MIDDLE,
        shape=hb,
    )
    text(s, x + 0.12, y + 0.72, 2.74, 1.15, [(b, 13, False, INK)])

# ---------------------------------------------------------------- 5. Module detail: agents
s = S[4]
frame(
    s,
    "Module – Multi-Agent Investigation",
    "The core module: what each agent does and the guardrails around it.",
)
text(
    s,
    0.45,
    1.4,
    7.4,
    0.4,
    [("Subtopic: six bounded agents and a deterministic engine", 18, True, BLUE)],
)
agents = [
    (
        "Triage",
        "Class (20-class taxonomy), domain SAFETY / QUALITY / BOTH, severity, priority, missing info",
    ),
    (
        "Safety investigation",
        "Unsafe acts, unsafe conditions, failed or missing controls",
    ),
    (
        "Quality investigation",
        "Defect, nonconformance and the inspection / test points that missed it",
    ),
    (
        "Risk engine (no AI)",
        "Likelihood × consequence on a 5×5 matrix → band; uses the higher of the two views",
    ),
    (
        "Root cause (RCA)",
        "5-Whys across people, process, equipment, materials, environment, management",
    ),
    (
        "Compliance",
        "Maps facts to retrieved procedure clauses: likely not met / unclear – never certifies",
    ),
    (
        "CAPA",
        "Corrective and preventive actions with owner, due date and verification criterion",
    ),
]
for i, (h, b) in enumerate(agents):
    y = 1.9 + i * 0.6
    fill = TEAL if h.startswith("Risk") else NAVY
    nb = box(s, 0.45, y, 2.35, 0.5, fill=fill, radius=0.2)
    text(
        s,
        0,
        0,
        0,
        0,
        [(h, 13, True, WHITE)],
        align=PP_ALIGN.CENTER,
        anchor=MSO_ANCHOR.MIDDLE,
        shape=nb,
        margin=0.03,
    )
    text(s, 2.9, y, 5.0, 0.5, [(b, 13, False, INK)], anchor=MSO_ANCHOR.MIDDLE)
g = box(s, 8.2, 1.45, 4.7, 4.95, fill=PALE_G)
text(
    s,
    8.4,
    1.55,
    4.3,
    4.8,
    [
        ("Guardrails", 18, True, GREEN),
        {
            "runs": [
                ("Explicit tool allowlist and call budget per agent", 14, False, INK)
            ],
            "bullet": True,
        },
        {
            "runs": [
                (
                    "Outputs validated against Pydantic schemas; invalid output falls back to rules",
                    14,
                    False,
                    INK,
                )
            ],
            "bullet": True,
        },
        {
            "runs": [
                (
                    "Cites only evidence retrieved in the same run; unknown citations removed",
                    14,
                    False,
                    INK,
                )
            ],
            "bullet": True,
        },
        {
            "runs": [
                (
                    "Documents and uploads treated as untrusted data; injection attempts flagged",
                    14,
                    False,
                    INK,
                )
            ],
            "bullet": True,
        },
        {
            "runs": [
                (
                    "Disagreements stored as conflicts for a person to decide",
                    14,
                    False,
                    INK,
                )
            ],
            "bullet": True,
        },
        {
            "runs": [("Agents never approve, verify or close", 14, True, INK)],
            "bullet": True,
        },
        {
            "runs": [
                (
                    "Providers: deterministic rules (offline, default) or Claude (optional)",
                    14,
                    False,
                    INK,
                )
            ],
            "bullet": True,
        },
    ],
)

# ---------------------------------------------------------------- 6. Problem statement
s = S[5]
frame(s, "Problem Statement", "Why this system is needed.")
ps = box(s, 0.45, 1.5, 5.2, 4.1, fill=NAVY)
text(
    s,
    0,
    0,
    0,
    0,
    [
        ("Problem", 18, True, "BFDBFE"),
        (
            "Construction organisations lack a single, trustworthy system that turns a reported safety or quality incident "
            "into an evidence-backed investigation, proportionate corrective actions and verified closure – quickly, "
            "consistently and with a complete audit trail.",
            17,
            False,
            WHITE,
        ),
    ],
    shape=ps,
    margin=0.25,
    anchor=MSO_ANCHOR.MIDDLE,
)
pains = [
    (
        "Fragmented evidence",
        "Reports, photos, procedures, NCRs and action logs live in different tools and inboxes.",
    ),
    (
        "Slow, inconsistent investigations",
        "Root causes and actions depend on who investigates; procedures are not cited.",
    ),
    (
        "Weak follow-through",
        "Overdue actions and unverified closures go unnoticed until the next incident.",
    ),
    (
        "Unsafe automation",
        "Generic AI can hallucinate, follow injected instructions or “close” an incident – "
        "unacceptable for safety.",
    ),
]
for i, (h, b) in enumerate(pains):
    card(
        s,
        5.95,
        1.45 + i * 1.22,
        6.95,
        1.12,
        h,
        b,
        size=13,
        head_size=15,
        badge=str(i + 1),
        badge_fill=SKY,
    )

# ---------------------------------------------------------------- 7. Methodology: architecture
s = S[6]
frame(
    s,
    "Methodology – System Architecture",
    "Layered architecture, from users to data. Arrows show the request "
    "path. Everything runs in one Docker image; PostgreSQL is the system of record.",
)
L = 0.45
W = 9.6


def band(y, h, label, color):
    box(s, L, y, W, h, fill=color, radius=0.06)
    text(s, L + 0.08, y, 1.3, h, [(label, 12, True, NAVY)], anchor=MSO_ANCHOR.MIDDLE)


band(1.42, 0.55, "Users", PALE)
node(
    s,
    1.85,
    1.49,
    8.1,
    0.41,
    "HSE manager · Safety & QA/QC engineers · Site engineer · Project manager · Auditor",
    None,
    BLUE,
    tsize=13,
)
band(2.12, 0.65, "Access", PALE)
node(s, 1.85, 2.19, 3.9, 0.51, "Web app", "Next.js · TypeScript · Tailwind", SKY)
node(
    s,
    6.05,
    2.19,
    3.9,
    0.51,
    "FastAPI REST API",
    "JWT · RBAC · tenant scoping · rate limits",
    SKY,
)
band(2.92, 0.75, "Services", PALE)
for i, (t, sub) in enumerate(
    (
        ("Incidents", "lifecycle"),
        ("Evidence", "files · vision"),
        ("CAPA", "escalation"),
        ("Reports", "MD · Word · PDF"),
        ("Audit", "notifications"),
    )
):
    node(s, 1.85 + i * 1.64, 2.99, 1.52, 0.61, t, sub, TEAL, ssize=10)
band(3.82, 1.0, "AI layer", PALE)
node(
    s,
    1.85,
    3.9,
    1.6,
    0.84,
    "Orchestrator",
    "checkpoints · pause / resume",
    NAVY,
    ssize=10,
)
for i, a in enumerate(("Triage", "Safety", "Quality", "RCA", "Compliance", "CAPA")):
    node(
        s,
        3.6 + (i % 3) * 1.12,
        3.9 + (i // 3) * 0.44,
        1.05,
        0.38,
        a,
        None,
        BLUE,
        tsize=11,
    )
node(s, 7.0, 3.9, 1.45, 0.84, "Risk engine", "5×5 matrix (no AI)", TEAL, ssize=10)
node(s, 8.55, 3.9, 1.4, 0.84, "Human gate", "approve · verify · close", GREEN, ssize=10)
band(4.97, 0.75, "Data", PALE)
x = 1.85
for t, sub, w in (
    ("PostgreSQL 16", "system of record", 1.7),
    ("pgvector + BM25", "hybrid RAG", 1.7),
    ("Neo4j", "knowledge graph", 1.4),
    ("Evidence store", "SHA-256", 1.45),
    ("LLM provider", "rules / Claude", 1.45),
):
    node(s, x, 5.04, w, 0.61, t, sub, NAVY, ssize=10)
    x += w + 0.1
for y1, y2 in ((1.97, 2.12), (2.77, 2.92), (3.67, 3.82), (4.82, 4.97)):
    arrow(s, 5.9, y1, 5.9, y2, color=NAVY, width=2)
side = box(s, 10.3, 1.42, 2.6, 4.3, fill=PALE_G)
text(
    s,
    10.38,
    1.5,
    2.45,
    4.2,
    [
        ("Platform", 15, True, GREEN),
        {"runs": [("One Docker image: API + web app", 12, False, INK)], "bullet": True},
        {"runs": [("Alembic migrations on start", 12, False, INK)], "bullet": True},
        {"runs": [("GitHub Actions: 7 CI jobs", 12, False, INK)], "bullet": True},
        {
            "runs": [
                (
                    "Azure Bicep: Container Apps, PostgreSQL, Key Vault, Files",
                    12,
                    False,
                    INK,
                )
            ],
            "bullet": True,
        },
        {
            "runs": [
                ("Security headers, CSP, login throttle, upload checks", 12, False, INK)
            ],
            "bullet": True,
        },
        {
            "runs": [("Structured logs with correlation IDs", 12, False, INK)],
            "bullet": True,
        },
    ],
)
text(
    s,
    2.4,
    5.85,
    10.5,
    0.35,
    [
        (
            "Request path: user → web app → API (authorisation enforced server-side) → services "
            "→ agents → data. The browser never decides access.",
            12,
            False,
            MUTED,
        )
    ],
)

# ---------------------------------------------------------------- 8. Methodology: workflow
s = S[7]
frame(
    s,
    "Methodology – Incident Resolution Workflow",
    "Read the top row left to right, then the bottom row right "
    "to left. Blue = AI agent, teal = deterministic, green = a person decides.",
)
row1 = [
    ("Report incident", "+ photos, documents", GREEN),
    ("Triage agent", "class · domain · severity", BLUE),
    ("Safety / Quality agents", "acts · conditions · controls", BLUE),
    ("Risk matrix", "5×5 → band", TEAL),
    ("RCA agent", "5-Whys", BLUE),
    ("Compliance agent", "cited clauses", BLUE),
]
w, gap = 1.83, 0.3
for i, (t, sub, c) in enumerate(row1):
    node(s, 0.45 + i * (w + gap), 1.55, w, 1.0, t, sub, c, tsize=13, ssize=11)
    if i:
        arrow(
            s,
            0.45 + i * (w + gap) - gap + 0.03,
            2.05,
            0.45 + i * (w + gap) - 0.03,
            2.05,
            color=NAVY,
        )
row2 = [
    ("CAPA agent", "corrective & preventive", BLUE),
    ("Human review", "approve / modify / reject", GREEN),
    ("Action work", "owners · due dates · escalation", GREEN),
    ("Verification", "independent person", GREEN),
    ("Closure", "authorised role", GREEN),
]
x_last = 0.45 + 5 * (w + gap)
for i, (t, sub, c) in enumerate(row2):
    x = x_last - i * (w + gap)
    node(s, x, 3.15, w, 1.0, t, sub, c, tsize=13, ssize=11)
    if i:
        arrow(s, x + w + gap - 0.03, 3.65, x + w + 0.03, 3.65, color=NAVY)
arrow(s, x_last + w / 2, 2.58, x_last + w / 2, 3.12, color=NAVY)
text(
    s,
    0.45 + 0 * (w + gap),
    3.2,
    w,
    0.9,
    [
        (
            "Checkpointed after every stage: pause, resume or cancel; crashed runs "
            "resume without repeating work.",
            12,
            False,
            MUTED,
        )
    ],
    anchor=MSO_ANCHOR.MIDDLE,
)
text(
    s, 2.4, 4.4, 10.5, 0.35, [("Incident status (enforced by the API)", 14, True, NAVY)]
)
states = [
    "Reported",
    "Under investigation",
    "Pending approval",
    "Action in progress",
    "Pending verification",
    "Closed",
]
sw = 1.55
for i, st in enumerate(states):
    x = 2.4 + i * (sw + 0.22)
    b = box(
        s,
        x,
        4.8,
        sw,
        0.5,
        fill=PALE if i < 5 else PALE_G,
        line=NAVY if i < 5 else GREEN,
        radius=0.3,
    )
    text(
        s,
        0,
        0,
        0,
        0,
        [(st, 12, True, NAVY if i < 5 else GREEN)],
        align=PP_ALIGN.CENTER,
        anchor=MSO_ANCHOR.MIDDLE,
        shape=b,
        margin=0.02,
    )
    if i:
        arrow(s, x - 0.2, 5.05, x - 0.02, 5.05, color=GREY, width=1.5)
text(
    s,
    2.4,
    5.45,
    10.5,
    0.75,
    [
        {
            "runs": [
                (
                    "Critical actions need the HSE manager; an assignee cannot verify their own action; an incident closes "
                    "only from Pending verification and only when every AI conflict has a recorded human decision.",
                    13,
                    False,
                    INK,
                )
            ]
        },
    ],
)

# ---------------------------------------------------------------- 9. Methodology: development & evaluation
s = S[8]
frame(
    s,
    "Methodology – Development & Evaluation",
    "How the system was built and how accuracy was measured without fooling ourselves.",
)
text(
    s,
    0.45,
    1.4,
    6.2,
    0.4,
    [("Staged delivery, every change through CI", 17, True, BLUE)],
)
text(
    s,
    0.45,
    1.8,
    6.2,
    0.35,
    [
        (
            "READ → INSPECT → PLAN → IMPLEMENT → TEST → REVIEW → DOCUMENT → COMMIT",
            11,
            True,
            MUTED,
        )
    ],
)
stages = [
    ("A", "Hardening, Azure IaC, staging, backup / restore, smoke test"),
    ("B", "Document parsing, embeddings, hybrid retrieval, grounded answers"),
    ("C", "Word / PDF investigation reports, email outbox"),
    ("D", "Vision analysis, knowledge graph, resumable workflows, conflicts"),
    ("E", "Web screens for every feature, browser tests"),
    ("+", "Accuracy round, load testing, concurrency fix"),
]
for i, (k, t) in enumerate(stages):
    y = 2.25 + i * 0.58
    c = box(
        s,
        0.6,
        y,
        0.45,
        0.45,
        fill=[NAVY, BLUE, SKY, TEAL, GREEN, NAVY][i],
        shape=MSO_SHAPE.OVAL,
    )
    text(
        s,
        0,
        0,
        0,
        0,
        [(k, 14, True, WHITE)],
        align=PP_ALIGN.CENTER,
        anchor=MSO_ANCHOR.MIDDLE,
        shape=c,
        margin=0,
    )
    text(s, 1.2, y - 0.02, 5.4, 0.5, [(t, 14, False, INK)], anchor=MSO_ANCHOR.MIDDLE)
ev = box(s, 6.95, 1.4, 5.95, 4.75, fill=PALE)
text(s, 7.15, 1.5, 5.6, 0.45, [("Honest evaluation protocol", 17, True, NAVY)])
steps = [
    (
        "1",
        "Freeze",
        "A separate author writes 46 held-out cases; committed before any tuning",
    ),
    (
        "2",
        "Tune on dev only",
        "Rules and procedures improved using 120 dev cases; held-out never opened",
    ),
    (
        "3",
        "Measure once",
        "Held-out scored once on the frozen rules and published as is",
    ),
    (
        "4",
        "Gate in CI",
        "Floors just under the results, 0 CRITICAL under-rated, 10/10 red-team",
    ),
]
for i, (k, h, t) in enumerate(steps):
    y = 2.05 + i * 1.0
    c = box(s, 7.15, y + 0.05, 0.45, 0.45, fill=GREEN, shape=MSO_SHAPE.OVAL)
    text(
        s,
        0,
        0,
        0,
        0,
        [(k, 14, True, WHITE)],
        align=PP_ALIGN.CENTER,
        anchor=MSO_ANCHOR.MIDDLE,
        shape=c,
        margin=0,
    )
    text(s, 7.75, y - 0.05, 5.0, 0.95, [(h, 15, True, INK), (t, 13, False, MUTED)])

# ---------------------------------------------------------------- 10. Approach / algorithms / tools
s = S[9]
frame(
    s,
    "Approach, Algorithms & Tools",
    "Key algorithms on the left, technology stack on the right.",
)
algos = [
    (
        "Weighted keyword triage",
        "20-class taxonomy; title words ×2, ambiguous words ×0.5; falling-object, fall and "
        "struck-by patterns",
    ),
    (
        "Hybrid retrieval (RAG)",
        "BM25 + vector search fused by Reciprocal Rank Fusion, then a deterministic rerank "
        "on coverage, headings and phrases",
    ),
    (
        "Deterministic risk scoring",
        "Likelihood (1–5) × consequence (1–5) → LOW / MEDIUM / HIGH / CRITICAL; the AI only "
        "suggests inputs",
    ),
    (
        "Root cause & grounding",
        "5-Whys over six cause categories; answers keep only claims supported by cited "
        "chunks",
    ),
]
for i, (h, b) in enumerate(algos):
    card(
        s,
        0.45,
        1.38 + i * 1.15,
        7.1,
        1.1,
        h,
        b,
        size=13,
        head_size=15,
        badge=str(i + 1),
        badge_fill=[NAVY, BLUE, TEAL, GREEN][i],
    )
text(s, 7.85, 1.4, 5.0, 0.4, [("Tools & technologies", 17, True, NAVY)])
tools = [
    ("Python 3.12", NAVY),
    ("FastAPI", NAVY),
    ("Pydantic", NAVY),
    ("SQLAlchemy", NAVY),
    ("Alembic", NAVY),
    ("PostgreSQL 16", TEAL),
    ("pgvector", TEAL),
    ("Neo4j 5", TEAL),
    ("Next.js", BLUE),
    ("TypeScript", BLUE),
    ("Tailwind CSS", BLUE),
    ("Docker", SKY),
    ("GitHub Actions", SKY),
    ("Azure Bicep", SKY),
    ("pytest", GREEN),
    ("Playwright", GREEN),
    ("ruff · mypy", GREEN),
    ("Claude (optional)", GREY),
]
for i, (t, c) in enumerate(tools):
    col, row = i % 3, i // 3
    b = box(s, 7.85 + col * 1.69, 1.9 + row * 0.66, 1.57, 0.52, fill=c, radius=0.3)
    text(
        s,
        0,
        0,
        0,
        0,
        [(t, 12, True, WHITE)],
        align=PP_ALIGN.CENTER,
        anchor=MSO_ANCHOR.MIDDLE,
        shape=b,
        margin=0.02,
    )

# ---------------------------------------------------------------- 11. Findings: accuracy
s = S[10]
frame(
    s,
    "Findings – AI Accuracy on Held-out Data",
    "One measurement before and one after tuning, on 46 held-out "
    "cases that were never used for tuning. Scores are for the offline rules provider on synthetic cases.",
)
cd = CategoryChartData()
cd.categories = [
    "Incident class",
    "Domain",
    "Severity ±1 level",
    "Right procedure (top 3)",
]
cd.add_series("Before tuning", (67.4, 80.4, 95.5, 67.4))
cd.add_series("After tuning", (89.1, 93.5, 97.7, 93.5))
cd.add_series("Target", (85, 90, 90, 85))
gf = s.shapes.add_chart(
    XL_CHART_TYPE.COLUMN_CLUSTERED,
    Inches(0.45),
    Inches(1.4),
    Inches(8.3),
    Inches(4.3),
    cd,
)
ch = gf.chart
ch.has_legend = True
ch.legend.position = XL_LEGEND_POSITION.TOP
ch.legend.include_in_layout = False
ch.legend.font.size = Pt(12)
ch.legend.font.color.rgb = rgb(INK)
ch.value_axis.maximum_scale = 100
ch.value_axis.minimum_scale = 0
ch.value_axis.has_major_gridlines = True
ch.value_axis.major_gridlines.format.line.color.rgb = rgb("E5E7EB")
ch.value_axis.tick_labels.font.size = Pt(11)
ch.value_axis.tick_labels.font.color.rgb = rgb(MUTED)
ch.value_axis.format.line.fill.background()
ch.category_axis.tick_labels.font.size = Pt(12)
ch.category_axis.tick_labels.font.color.rgb = rgb(INK)
plot = ch.plots[0]
plot.gap_width = 60
plot.overlap = -5
for ser, col in zip(plot.series, ("9DB7D5", NAVY, GREEN), strict=True):
    ser.format.fill.solid()
    ser.format.fill.fore_color.rgb = rgb(col)
    ser.data_labels.show_value = True
    ser.data_labels.number_format = '0.0"%"'
    ser.data_labels.number_format_is_linked = False
    ser.data_labels.position = XL_LABEL_POSITION.OUTSIDE_END
    ser.data_labels.font.size = Pt(10)
    ser.data_labels.font.color.rgb = rgb(INK)
stats = [
    ("0", "CRITICAL incidents rated LOW or MEDIUM"),
    ("10 / 10", "red-team attacks contained"),
    ("100%", "workflow invariants held"),
    ("166", "labelled cases: 120 dev + 46 held-out"),
]
for i, (big, small) in enumerate(stats):
    b = box(s, 9.05, 1.45 + i * 1.12, 3.85, 1.0, fill=PALE_G if i < 3 else PALE)
    text(
        s,
        9.2,
        1.45 + i * 1.12,
        1.4,
        1.0,
        [(big, 24, True, GREEN if i < 3 else NAVY)],
        anchor=MSO_ANCHOR.MIDDLE,
    )
    text(
        s,
        10.55,
        1.45 + i * 1.12,
        2.3,
        1.0,
        [(small, 13, False, INK)],
        anchor=MSO_ANCHOR.MIDDLE,
    )
text(
    s,
    2.4,
    5.85,
    10.5,
    0.35,
    [
        (
            "All four accuracy targets met; class +21.7 points and retrieval +26.1 points "
            "after tuning.",
            13,
            True,
            NAVY,
        )
    ],
)

# ---------------------------------------------------------------- 12. Findings: end-to-end resolution
s = S[11]
frame(
    s,
    "Findings – End-to-End Incident Resolution",
    "resolve_all.py drove the running product over HTTP on "
    "PostgreSQL: the agents investigated, scripted reviewers decided, and every incident closed.",
)
rows = [
    ("Metric", "Before", "After"),
    ("Incidents", "20", "21"),
    ("Open incidents", "13", "0"),
    ("Actions awaiting review", "15", "0"),
    ("Actions overdue", "4", "0"),
    ("Mean days to close*", "11", "6.1"),
]
tbl = s.shapes.add_table(
    len(rows), 3, Inches(0.45), Inches(1.45), Inches(5.0), Inches(2.9)
).table
tbl.columns[0].width = Inches(2.8)
tbl.columns[1].width = Inches(1.1)
tbl.columns[2].width = Inches(1.1)
for r, row in enumerate(rows):
    for c, v in enumerate(row):
        cell = tbl.cell(r, c)
        cell.text = v
        para = cell.text_frame.paragraphs[0]
        para.alignment = PP_ALIGN.LEFT if c == 0 else PP_ALIGN.CENTER
        f = para.runs[0].font
        f.size = Pt(14)
        f.name = "Calibri"
        f.bold = r == 0 or c == 2
        f.color.rgb = rgb(WHITE if r == 0 else (GREEN if c == 2 else INK))
        cell.fill.solid()
        cell.fill.fore_color.rgb = rgb(NAVY if r == 0 else (PALE if r % 2 else WHITE))
for i, (big, small) in enumerate(
    (
        ("107", "agent runs, 0 failures"),
        ("79", "human decisions: 64 approved, 15 rejected"),
        ("21 / 21", "incidents closed after verification"),
    )
):
    b = box(s, 0.45 + i * 1.7, 4.55, 1.6, 1.15, fill=PALE_G)
    text(
        s,
        0,
        0,
        0,
        0,
        [(big, 20, True, GREEN), (small, 10, False, INK)],
        align=PP_ALIGN.CENTER,
        anchor=MSO_ANCHOR.MIDDLE,
        shape=b,
        margin=0.04,
    )
s.shapes.add_picture(
    "img/dashboard_after.png", Inches(5.8), Inches(1.45), width=Inches(7.1)
)
text(
    s,
    5.8,
    5.5,
    7.1,
    0.6,
    [
        (
            "Dashboard after the run: 21 incidents closed, 0 open, 0 overdue. *Scripted reviewers "
            "act within minutes, so days-to-close is a workflow check, not a performance "
            "figure.",
            11,
            False,
            MUTED,
        )
    ],
)


# ---------------------------------------------------------------- 13a/13b. Findings: the product in use
def shot(slide, path, x, y, w, caption, px):
    """Screenshot scaled to width w with a hairline border and a caption underneath; returns its bottom edge."""
    pic = slide.shapes.add_picture(path, Inches(x), Inches(y), width=Inches(w))
    pic.line.color.rgb = rgb("D1D5DB")
    pic.line.width = Pt(0.75)
    h = w * px[1] / px[0]
    text(slide, x, y + h + 0.02, w, 0.3, [(caption, 11, False, MUTED)])
    return y + h + 0.32


s = SCREENS[0]
frame(
    s,
    "Findings – The Product: Investigation to Closure",
    "Screens from the running product (browser end-to-end "
    "run on main). AI output is labelled; actions wait for the HSE manager; closure needs verified actions.",
)
shot(
    s,
    "img/s_workspace.png",
    0.45,
    1.4,
    5.6,
    "AI investigation of a new report: triage, safety findings, 5×5 risk, "
    "root causes, all marked AI-generated",
    (1280, 900),
)
y = shot(
    s,
    "img/s_queue.png",
    6.3,
    1.4,
    6.6,
    "Action queue: every AI-proposed action waits for a person; critical "
    "ones only for the HSE manager",
    (1280, 330),
)
shot(
    s,
    "img/s_closed.png",
    6.3,
    y + 0.08,
    6.6,
    "The same incident closed, only after each approved action was verified effective",
    (1280, 440),
)

s = SCREENS[1]
frame(
    s,
    "Findings – The Product: Knowledge, Photos & Audit",
    "Grounded answers cite the procedure passage they "
    "come from; photo findings are suggestions for an engineer; every action is in the audit trail.",
)
y = shot(
    s,
    "img/s_answer.png",
    0.45,
    1.4,
    6.25,
    "Grounded answer: each statement cites its source passage, or "
    "says the documents do not answer",
    (1280, 375),
)
shot(
    s,
    "img/s_audit.png",
    0.45,
    y + 0.1,
    6.25,
    "Audit trail: every human and agent action with a correlation ID",
    (1280, 420),
)
y = shot(
    s,
    "img/s_photo.png",
    6.95,
    1.4,
    5.95,
    "Photo review: image checks for an engineer to confirm; never proof on their own",
    (905, 350),
)
kb = box(s, 6.95, y + 0.12, 5.95, 5.75 - y - 0.12, fill=PALE)
text(
    s,
    0,
    0,
    0,
    0,
    [
        ("What the screens show", 15, True, NAVY),
        {
            "runs": [
                (
                    "AI output is always labelled and never approves, verifies or closes",
                    12,
                    False,
                    INK,
                )
            ],
            "bullet": True,
        },
        {
            "runs": [
                (
                    "Answers and requirements cite retrieved passages only",
                    12,
                    False,
                    INK,
                )
            ],
            "bullet": True,
        },
        {
            "runs": [
                (
                    "Role checks happen in the API; the browser only renders what it is allowed",
                    12,
                    False,
                    INK,
                )
            ],
            "bullet": True,
        },
    ],
    shape=kb,
    margin=0.18,
    anchor=MSO_ANCHOR.MIDDLE,
)

# ---------------------------------------------------------------- 13. Findings: performance & quality
s = S[12]
frame(
    s,
    "Findings – Performance & Engineering Quality",
    "Load test of the production image on PostgreSQL 16 "
    "(one API process, 4 vCPU host) and the quality gates that run on every change.",
)
lrows = [
    ("Load test run", "Users", "Req/s", "Errors", "Read p95"),
    ("Single user", "1", "53", "0%", "28 ms"),
    ("Saturation", "25", "56", "0%", "745 ms"),
    ("Realistic mix (2 s think time)", "100 + 2", "39", "0%", "802 ms"),
    ("8 concurrent reporters", "1 + 8", "11", "0%", "718 ms"),
]
tbl = s.shapes.add_table(
    len(lrows), 5, Inches(0.45), Inches(1.45), Inches(7.4), Inches(2.5)
).table
for c, wdt in enumerate((3.1, 1.0, 0.9, 1.0, 1.4)):
    tbl.columns[c].width = Inches(wdt)
for r, row in enumerate(lrows):
    for c, v in enumerate(row):
        cell = tbl.cell(r, c)
        cell.text = v
        para = cell.text_frame.paragraphs[0]
        para.alignment = PP_ALIGN.LEFT if c == 0 else PP_ALIGN.CENTER
        f = para.runs[0].font
        f.size = Pt(13)
        f.name = "Calibri"
        f.bold = r == 0
        f.color.rgb = rgb(WHITE if r == 0 else INK)
        cell.fill.solid()
        cell.fill.fore_color.rgb = rgb(NAVY if r == 0 else (PALE if r % 2 else WHITE))
fb = box(s, 0.45, 4.15, 7.4, 1.45, fill=PALE_G)
text(
    s,
    0.6,
    4.2,
    7.1,
    1.4,
    [
        ("Bug found and fixed by the load test", 15, True, GREEN),
        (
            "Two reports submitted at the same moment received the same reference and one failed (HTTP 500). Fixed with "
            "a PostgreSQL advisory lock plus a retry; 113 simultaneous reports later, no collision.",
            13,
            False,
            INK,
        ),
    ],
)
gates = [
    ("270", "backend tests (unit, API, migrations, evaluation, red-team)"),
    ("6", "browser end-to-end tests"),
    ("7", "CI jobs green on every pull request"),
    ("0", "open dependency-audit findings (Python + npm)"),
]
for i, (big, small) in enumerate(gates):
    b = box(s, 8.15, 1.45 + i * 1.07, 4.75, 0.95, fill=PALE)
    text(
        s,
        8.3,
        1.45 + i * 1.07,
        1.1,
        0.95,
        [(big, 24, True, NAVY)],
        anchor=MSO_ANCHOR.MIDDLE,
    )
    text(
        s,
        9.4,
        1.45 + i * 1.07,
        3.45,
        0.95,
        [(small, 13, False, INK)],
        anchor=MSO_ANCHOR.MIDDLE,
    )
text(
    s,
    2.4,
    5.85,
    10.5,
    0.35,
    [
        (
            "One process serves ≈ 55 req/s – enough for a pilot with tens of active users; the "
            "dashboard is the first thing to optimise before scaling.",
            13,
            True,
            NAVY,
        )
    ],
)

# ---------------------------------------------------------------- 14. Conclusion
s = S[13]
frame(
    s,
    "Conclusion",
    "What was achieved, what is honestly still open, and the next steps.",
)
text(s, 0.45, 1.4, 6.2, 0.4, [("Achieved", 18, True, GREEN)])
done = [
    "End-to-end platform: report → AI investigation → human approval → verified closure, with a full audit trail",
    "Six bounded agents plus a deterministic risk engine; outputs validated, cited and never self-approving",
    "Accuracy targets met on held-out data: class 89%, domain 93%, retrieval 93%, 0 critical under-rated",
    "Production engineering: secure API, 7-job CI, Docker image, Azure templates, load-tested",
]
for i, t in enumerate(done):
    card_y = 1.85 + i * 0.98
    b = box(s, 0.45, card_y, 6.2, 0.86, fill=PALE_G)
    c = box(s, 0.58, card_y + 0.2, 0.45, 0.45, fill=GREEN, shape=MSO_SHAPE.OVAL)
    text(
        s,
        0,
        0,
        0,
        0,
        [("✓", 14, True, WHITE)],
        align=PP_ALIGN.CENTER,
        anchor=MSO_ANCHOR.MIDDLE,
        shape=c,
        margin=0,
    )
    text(s, 1.15, card_y, 5.4, 0.86, [(t, 13, False, INK)], anchor=MSO_ANCHOR.MIDDLE)
text(s, 6.95, 1.4, 5.95, 0.4, [("Limitations & future work", 18, True, NAVY)])
todo = [
    (
        "Real data",
        "Expert-labelled pilot reports; real OSHA Severe Injury Reports scoring (pipeline ready)",
    ),
    (
        "Claude provider",
        "Score the LLM provider on the same held-out protocol with an API key",
    ),
    (
        "Scale",
        "SQL aggregates and pagination for the dashboard; shared login throttle before scaling out",
    ),
    (
        "Deploy & extend",
        "Azure staging; licensed vision model; graph-aware retrieval; OIDC sign-in",
    ),
]
for i, (h, t) in enumerate(todo):
    y = 1.85 + i * 0.98
    box(s, 6.95, y, 5.95, 0.86, fill=PALE)
    text(
        s,
        7.1,
        y,
        5.7,
        0.86,
        [((h + ": ", 13, True, NAVY), (t, 13, False, INK))],
        anchor=MSO_ANCHOR.MIDDLE,
    )
text(
    s,
    2.4,
    5.85,
    10.5,
    0.35,
    [
        (
            "Site Guard AI is decision support: people stay accountable for every safety "
            "decision, and the system never certifies legal compliance.",
            13,
            True,
            NAVY,
        )
    ],
)

# ---------------------------------------------------------------- 15. Acknowledgement
s = S[14]
frame(s, "Acknowledgement")
ab = box(s, 1.8, 1.7, 9.7, 3.6, fill=PALE)
text(
    s,
    0,
    0,
    0,
    0,
    [
        (
            "I sincerely thank the University and LTIMindtree for the Advanced Industry Integrated M.Tech / MCA Program, and "
            "L&T EduTech as implementation partner, for the opportunity and guidance to carry out this project.",
            18,
            False,
            INK,
        ),
        ("", 10, False, INK),
        (
            "I am grateful to my project guide, mentors and reviewers for their valuable feedback across the reviews, and to "
            "my family and peers for their constant support.",
            18,
            False,
            INK,
        ),
        ("", 10, False, INK),
        ("– S.K. Sricharan (T9984)", 18, True, NAVY),
    ],
    shape=ab,
    margin=0.35,
    anchor=MSO_ANCHOR.MIDDLE,
)

# ---------------------------------------------------------------- 16. References
s = S[15]
frame(s, "References")
refs = [
    "U.S. OSHA, 29 CFR Part 1926 – Safety and Health Regulations for Construction (Subparts E, K, L, M, P, CC).",
    "U.S. OSHA, Severe Injury Reports dataset, https://www.osha.gov/severeinjury",
    "UK: The Work at Height Regulations 2005; CDM Regulations 2015; LOLER 1998; RIDDOR 2013.",
    "India: Building and Other Construction Workers Act 1996; OSH and Working Conditions Code 2020.",
    "ISO 45001:2018, Occupational health and safety management systems – Requirements.",
    "S. Robertson and H. Zaragoza, “The Probabilistic Relevance Framework: BM25 and Beyond,” Foundations and Trends "
    "in Information Retrieval, 2009.",
    "G. V. Cormack, C. L. A. Clarke and S. Büttcher, “Reciprocal Rank Fusion outperforms Condorcet and individual "
    "rank learning methods,” SIGIR 2009.",
    "P. Lewis et al., “Retrieval-Augmented Generation for Knowledge-Intensive NLP Tasks,” NeurIPS 2020.",
    "OWASP, Top 10 for Large Language Model Applications (prompt injection, insecure output handling).",
    "Documentation: FastAPI, SQLAlchemy, PostgreSQL / pgvector, Neo4j, Next.js, Playwright, Anthropic Claude API.",
    "Project repository: github.com/sricharansk/SITEGUARD-AI (docs/EVALUATION.md, docs/LOAD_TEST.md, docs/RESULTS.md).",
]
text(
    s,
    0.45,
    1.45,
    12.45,
    4.4,
    [
        {"runs": [(f"[{i + 1}]  ", 15, True, NAVY), (r, 15, False, INK)], "space": 7}
        for i, r in enumerate(refs)
    ],
)

prs.save(OUT)
print("saved", OUT)
